"""Read an uploaded requirements document (txt, md or docx) and return plain text.

The text goes to detect_drift(), which splits it into sentences. So this module removes things that are
not requirements and would show up as fake changes: markdown headings, bold marks, file titles.
"""
import io
import re
from pathlib import Path

SUPPORTED = {".txt", ".md", ".docx"}
MAX_BYTES = 1_000_000  # 1 MB, same limit as code files


class DocumentError(ValueError):
    """The file cannot be used. The message is safe to show to the user."""


def clean_markdown(text: str) -> str:
    """Drop heading lines (# Title), horizontal rules and bold/code marks. Keep the requirement lines."""
    kept = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#") or re.fullmatch(r"[-*_=]{3,}", stripped):
            continue
        line = line.replace("**", "").replace("__", "").replace("`", "")
        kept.append(line.rstrip())
    return "\n".join(kept)


def _read_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise DocumentError("Word files need the python-docx package. Use .txt or .md instead.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("This .docx file could not be opened.") from exc

    lines = []
    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style.startswith(("Heading", "Title")):
            continue
        lines.append(paragraph.text)
    for table in document.tables:  # requirements are often written in tables
        for row in table.rows:
            lines.append(" ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(lines)


def read_document(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"Unsupported file type '{ext or filename}'. Use .txt, .md or .docx.")
    if len(data) > MAX_BYTES:
        raise DocumentError("File is too large (limit 1 MB).")

    if ext == ".docx":
        text = _read_docx(data)
    else:
        text = data.decode("utf-8-sig", errors="replace")  # utf-8-sig drops the BOM that Notepad adds
        if ext == ".md":
            text = clean_markdown(text)

    text = text.strip()
    if not text:
        raise DocumentError("The document has no text.")
    return text
